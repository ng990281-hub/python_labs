import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """
    Конвертирует CSV в XLSX.
    Использовать openpyxl ИЛИ xlsxwriter.
    Первая строка CSV — заголовок.
    Лист называется "Sheet1".
    Колонки — автоширина по длине текста (не менее 8 символов).
    """
    cp=Path(csv_path)
    xlp=Path(xlsx_path)
    
    if not cp.is_file():
        raise FileNotFoundError("Файл не найден")
    if cp.suffix.lower() != ".csv":
        raise ValueError("Входной файл должен быть CSV")

    rows=[]
    with open(cp, encoding="utf-8", newline="") as f:
        for row in csv.reader(f):
            if row:
                rows.append(row)
    
    if len(rows) < 2:
        raise ValueError("CSV пустой или в нём нет данных (только заголовок)")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in rows:
        ws.append(row)
    wid = {}
    for row in rows:
        for i, value in enumerate(row, start=1):
            wid[i] = max(wid.get(i,0), len(value))
    for index, width in wid.items():
        letter = get_column_letter(index)
        ws.column_dimensions[letter].width = max(8, width)


    wb.save(xlsx_path)


if __name__ == "__main__":
    csv_to_xlsx("data/samples/people.csv", "data/out/people.xlsx")
    csv_to_xlsx("data/samples/cities.csv", "data/out/cities.xlsx")







        
